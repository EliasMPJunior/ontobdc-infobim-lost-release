from __future__ import annotations

import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from ontobdc.cli.domain.port.context import CliContextPort
from ontobdc.shared.adapter.capability import TransformationCapability
from ontobdc.shared.domain.model.capability import CapabilityMetadata
from ontobdc_view.page.adapter.context import PAGE_ELEMENT_URI, PageDataContextAdapter
from ontobdc_view.page.plugin.builder.ifc_work_schedule.gantt_payload import (
    GanttPayloadAdapter,
)
from ontobdc_view.shared.adapter.vendor import VENDOR_SHEET_JS_NAME
from ontobdc_view.surface.plugin.capability.transformation.data_gathered import (
    DataGatheredCapability,
)


_IBIM = "https://infobim.org/ontology/ns#"
_GANTT_ENTITY_NAMES: Sequence[str] = (
    "IfcWorkSchedule",
    "IfcTask",
    "IfcTaskTime",
    "IfcRelSequence",
)
_GANTT_DETAIL_ENTITY_NAMES: Sequence[str] = (
    "IfcTask",
    "IfcTaskTime",
    "IfcRelSequence",
)


class IfcWorkScheduleGanttPayloadGatheredCapability(TransformationCapability):
    """Build the runtime payload and initial graph for the IfcWorkSchedule Gantt.

    The stable browser Gantt reads four workbook sheets. During Page generation
    this capability materializes the task/time/sequence rows into the embedded
    JSON-LD graph so the generated Page can draw immediately. The browser-side
    folder connection remains available for live refresh and write-back.
    """

    METADATA = CapabilityMetadata(
        id=(
            "org.ontobdc.view.plugin.capability.transformation.target."
            "ifc_work_schedule_gantt_payload_gathered"
        ),
        version="1.0.0",
        name="IfcWorkSchedule Gantt Payload Gathered",
        description=(
            "Build the stable Gantt runtime context and materialize its "
            "task graph from the IfcWorkSchedule workbook."
        ),
        author=["http://kb.elias.eng.br/nid/elias.ttl#Elias"],
        tags=["view", "page", "ifc-work-schedule", "gantt", "transformation"],
        supported_languages=["en", "pt-br"],
        log_message={
            "info": {"en": "The IfcWorkSchedule Gantt payload was gathered."},
            "debug_entry": {
                "en": "Building the IfcWorkSchedule Gantt runtime payload and task graph."
            },
        },
    )

    SCRIPT_NAMES = [
        VENDOR_SHEET_JS_NAME,
        "i18n_apply",
        "graph_reader",
        "container_connection",
        "connection_state",
        "chrome_controls",
        "pyodide_runtime",
        "task_table_timeline",
        "dependency_arrows",
        "gantt_tab_mount",
    ]

    def label(self, lang: str = "en") -> str:
        return "IfcWorkSchedule Gantt Payload Gathered"

    def description(self, lang: str = "en") -> str:
        return self.METADATA.description

    def check(self, context: CliContextPort) -> bool:
        payload = PageDataContextAdapter.payload(context)
        return (
            isinstance(payload.get("@graph"), list)
            and "gantt_payload" in payload
            and isinstance(payload.get("gantt_script_names"), list)
        )

    def execute(self, context: CliContextPort) -> Dict[str, Any]:
        payload = PageDataContextAdapter.payload(context)
        source_node = PageDataContextAdapter.source_node(context)
        element_uri = PageDataContextAdapter.require_uri(context, PAGE_ELEMENT_URI)
        identifier = self._identifier(source_node, element_uri)

        payload.update(source_node)
        payload["@graph"] = self._source_graph(
            context,
            source_node=source_node,
            element_uri=element_uri,
        )
        payload["gantt_payload"] = GanttPayloadAdapter().build(
            source_node,
            element_uri,
            identifier,
        )
        payload["gantt_script_names"] = list(self.SCRIPT_NAMES)
        return {"resulting_state": "__gantt_payload_gathered__"}

    def is_satisfied(self, context: CliContextPort) -> bool:
        return self.check(context)

    @classmethod
    def _source_graph(
        cls,
        context: CliContextPort,
        *,
        source_node: Dict[str, Any],
        element_uri: str,
    ) -> List[Dict[str, Any]]:
        nodes = cls._data_gathered_nodes(context)
        source_id = str(source_node.get("@id") or "").strip()

        ordered: List[Dict[str, Any]] = [dict(source_node)]
        for node in nodes:
            node_id = str(node.get("@id") or "").strip()
            if source_id and node_id == source_id:
                continue
            ordered.append(node)

        # DATA_GATHERED currently materializes registered DataEntity rows, but
        # the schedule's IfcTask/IfcTaskTime/IfcRelSequence sheets are not
        # necessarily registered as independent dataset entities. Preserve any
        # task graph already present; otherwise derive exactly the same node
        # shape used by the stable browser SheetJS parser.
        if not cls._contains_gantt_detail_nodes(ordered):
            ordered.extend(
                cls._gantt_detail_nodes_from_workbook(
                    context,
                    schedule_uri=element_uri,
                )
            )
        return ordered

    @staticmethod
    def _data_gathered_nodes(context: CliContextPort) -> List[Dict[str, Any]]:
        try:
            raw: Any = json.loads(
                DataGatheredCapability.state_path(context).read_text(encoding="utf-8")
            )
        except (OSError, RuntimeError, TypeError, ValueError, json.JSONDecodeError):
            return []

        if isinstance(raw, list):
            return [node for node in raw if isinstance(node, dict)]
        if isinstance(raw, dict) and isinstance(raw.get("@graph"), list):
            return [node for node in raw["@graph"] if isinstance(node, dict)]
        if isinstance(raw, dict):
            return [raw]
        return []

    @classmethod
    def _contains_gantt_detail_nodes(cls, nodes: Iterable[Dict[str, Any]]) -> bool:
        wanted = {_IBIM + name for name in _GANTT_DETAIL_ENTITY_NAMES}
        for node in nodes:
            raw_types = node.get("@type")
            types = raw_types if isinstance(raw_types, list) else [raw_types]
            if any(str(value or "") in wanted for value in types):
                return True
        return False

    @classmethod
    def _gantt_detail_nodes_from_workbook(
        cls,
        context: CliContextPort,
        *,
        schedule_uri: str,
    ) -> List[Dict[str, Any]]:
        dataset_path = cls._dataset_path(context, schedule_uri)
        if dataset_path is None:
            return []
        workbook_path = cls._workbook_path(dataset_path)
        if workbook_path is None:
            return []

        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                filename=str(workbook_path),
                read_only=True,
                data_only=True,
            )
        except Exception:
            return []

        nodes: List[Dict[str, Any]] = []
        try:
            for entity_name in _GANTT_DETAIL_ENTITY_NAMES:
                if entity_name not in workbook.sheetnames:
                    continue
                rows = cls._sheet_records(workbook[entity_name])
                nodes.extend(
                    cls._records_to_nodes(
                        rows,
                        entity_name=entity_name,
                        schedule_uri=schedule_uri,
                    )
                )
        finally:
            try:
                workbook.close()
            except Exception:
                pass
        return nodes

    @staticmethod
    def _dataset_path(
        context: CliContextPort,
        element_uri: str,
    ) -> Optional[Path]:
        container_path = str(
            context.get_parameter_value("container_path") or ""
        ).strip()
        if not container_path:
            return None
        segments = [segment for segment in str(element_uri or "").split("/") if segment]
        if len(segments) < 2:
            return None
        return Path(container_path).expanduser().resolve() / segments[-2]

    @classmethod
    def _workbook_path(cls, dataset_path: Path) -> Optional[Path]:
        datapackage_path = dataset_path / ".__ontobdc__" / "datapackage.json"
        candidates: List[Path] = []
        if datapackage_path.is_file():
            try:
                descriptor = json.loads(datapackage_path.read_text(encoding="utf-8"))
            except (OSError, ValueError, json.JSONDecodeError):
                descriptor = {}
            for resource in list(descriptor.get("resources") or []):
                if not isinstance(resource, dict):
                    continue
                if not cls._resource_is_schedule(resource):
                    continue
                for raw_path in cls._resource_paths(resource.get("path")):
                    resolved = (datapackage_path.parent / raw_path).resolve()
                    if resolved.suffix.lower() == ".xlsx" and resolved.is_file():
                        candidates.append(resolved)

        if not candidates:
            candidates = sorted(
                path
                for path in dataset_path.rglob("*.xlsx")
                if path.is_file() and not path.name.startswith("~$")
            )
        if not candidates:
            return None

        # Match the stable SheetJS behavior: prefer the workbook containing the
        # greatest number of Gantt sheets, then deterministic path order.
        scored = sorted(
            ((cls._workbook_score(path), str(path), path) for path in candidates),
            key=lambda item: (-item[0], item[1]),
        )
        return scored[0][2]

    @staticmethod
    def _resource_is_schedule(resource: Dict[str, Any]) -> bool:
        values = (
            resource.get("entityIdentifier"),
            resource.get("entityUri"),
            resource.get("name"),
        )
        for value in values:
            text = str(value or "").strip().lower().replace("-", "_")
            local = text.rsplit("#", 1)[-1].rstrip("/").rsplit("/", 1)[-1]
            if local in {"ifcworkschedule", "ifc_work_schedule"}:
                return True
        return False

    @staticmethod
    def _resource_paths(value: Any) -> List[Path]:
        values = value if isinstance(value, list) else [value]
        result: List[Path] = []
        for raw in values:
            text = str(raw or "").strip()
            if text and "://" not in text:
                result.append(Path(text))
        return result

    @staticmethod
    def _workbook_score(path: Path) -> int:
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                filename=str(path),
                read_only=True,
                data_only=True,
            )
            try:
                names = set(workbook.sheetnames)
                return sum(1 for name in _GANTT_ENTITY_NAMES if name in names)
            finally:
                workbook.close()
        except Exception:
            return -1

    @classmethod
    def _sheet_records(cls, sheet: Any) -> List[Dict[str, str]]:
        rows = sheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration:
            return []
        headers = [str(value or "").strip() for value in header_values]
        if not any(headers):
            return []

        records: List[Dict[str, str]] = []
        for values in rows:
            if not any(value not in (None, "") for value in values):
                continue
            record: Dict[str, str] = {}
            for index, header in enumerate(headers):
                if not header or index >= len(values):
                    continue
                text = cls._cell_text(values[index])
                if text:
                    record[header] = text
            if record:
                records.append(record)
        return records

    @staticmethod
    def _cell_text(value: Any) -> str:
        if value is None or value == "":
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, datetime):
            return value.isoformat(timespec="seconds")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, time):
            return value.isoformat(timespec="seconds")
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @classmethod
    def _records_to_nodes(
        cls,
        records: Sequence[Dict[str, str]],
        *,
        entity_name: str,
        schedule_uri: str,
    ) -> List[Dict[str, Any]]:
        nodes: List[Dict[str, Any]] = []
        for index, record in enumerate(records):
            global_id = str(record.get("GlobalId") or f"{entity_name}-{index}").strip()
            node: Dict[str, Any] = {
                "@id": f"{schedule_uri}/{entity_name}/{global_id}",
                "@type": [_IBIM + entity_name],
            }
            for column, value in record.items():
                if not column or value == "":
                    continue
                node[_IBIM + column] = [{"@value": value}]
            nodes.append(node)
        return nodes

    @staticmethod
    def _identifier(source_node: Dict[str, Any], element_uri: str) -> str:
        values = source_node.get("http://purl.org/dc/terms/identifier")
        picked = values[0] if isinstance(values, list) and values else values
        if isinstance(picked, dict):
            value = picked.get("@value") or picked.get("@id")
            if value is not None and str(value).strip():
                return str(value).strip()
        elif picked is not None and str(picked).strip():
            return str(picked).strip()

        value = str(element_uri or "").rstrip("/")
        if "#" in value:
            return value.rsplit("#", 1)[-1]
        if "/" in value:
            return value.rsplit("/", 1)[-1]
        return value
