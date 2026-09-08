import json

import pytest
from openpyxl import Workbook, load_workbook

from infobim.element.adapter.parameter import ElementParameterRepository

GLOBAL_ID_URI = "https://example.test/GlobalId"
NAME_URI = "https://example.test/Name"


def _project(tmp_path, rows):
    project = tmp_path / "project"
    dataset = project / "elements"
    marker = dataset / ".__ontobdc__"
    linkset = marker / "linkset"
    payload = dataset / "payload" / "document"
    project_marker = project / ".__ontobdc__"
    linkset.mkdir(parents=True)
    payload.mkdir(parents=True)
    project_marker.mkdir(parents=True)

    (project_marker / "container.ttl").write_text(
        """
@prefix obdc: <http://ontobdc.org/ontology/domain/ontobdc/ns.ttl#> .
@prefix prov: <http://www.w3.org/ns/prov#> .
<urn:project> a obdc:DataContainer ; obdc:hasEntityDataset <urn:dataset> .
<urn:dataset> a obdc:EntityDataset ; prov:atLocation "elements" .
""",
        encoding="utf-8",
    )
    (linkset / "facade.ttl").write_text(
        f"""
@prefix obdc: <http://ontobdc.org/ontology/domain/ontobdc/ns.ttl#> .
@prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
<urn:field:global-id> obdc:identifier "GlobalId" ;
    obdc:mapsToProperty <{GLOBAL_ID_URI}> ;
    obdc:fieldDatatype xsd:string ;
    obdc:isIdentifierField true .
<urn:field:name> obdc:identifier "Name" ;
    obdc:mapsToProperty <{NAME_URI}> ;
    obdc:fieldDatatype xsd:string ;
    obdc:isIdentifierField false .
""",
        encoding="utf-8",
    )

    workbook_path = payload / "elements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Element"
    sheet.append(["GlobalId", "Name"])
    for row in rows:
        sheet.append(row)
    workbook.save(workbook_path)
    workbook.close()

    (marker / "datapackage.json").write_text(
        json.dumps(
            {
                "resources": [
                    {
                        "name": "Element",
                        "path": "../payload/document/elements.xlsx",
                        "format": "xlsx",
                        "dialect": {"excel": {"sheet": "Element"}},
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    return project, workbook_path


def _values(workbook_path):
    workbook = load_workbook(workbook_path)
    try:
        return [tuple(row) for row in workbook["Element"].iter_rows(values_only=True)]
    finally:
        workbook.close()


def test_set_one_parameter_without_assuming_ifc(tmp_path):
    project, workbook_path = _project(tmp_path, [("E-1", "Old")])

    mutations = ElementParameterRepository(str(project)).set("E-1", NAME_URI, "New")

    assert len(mutations) == 1
    assert _values(workbook_path)[1] == ("E-1", "New")


def test_set_all_updates_every_occurrence_and_single_set_refuses_ambiguity(tmp_path):
    project, workbook_path = _project(
        tmp_path,
        [("E-1", "First"), ("E-1", "Second")],
    )
    repository = ElementParameterRepository(str(project))

    with pytest.raises(ValueError, match="use --all"):
        repository.set("E-1", NAME_URI, "Updated")

    mutations = repository.set(
        "E-1",
        NAME_URI,
        "Updated",
        all_occurrences=True,
    )
    assert len(mutations) == 2
    assert _values(workbook_path)[1:] == [
        ("E-1", "Updated"),
        ("E-1", "Updated"),
    ]


def test_unset_one_parameter(tmp_path):
    project, workbook_path = _project(tmp_path, [("E-1", "Old")])

    mutation = ElementParameterRepository(str(project)).unset("E-1", NAME_URI)

    assert mutation.old_value == "Old"
    assert mutation.new_value is None
    assert _values(workbook_path)[1] == ("E-1", None)


def test_unset_all_preserves_the_global_id_identifier(tmp_path):
    project, workbook_path = _project(tmp_path, [("E-1", "Old")])

    mutations = ElementParameterRepository(str(project)).unset_all("E-1")

    assert [mutation.parameter_uri for mutation in mutations] == [NAME_URI]
    assert _values(workbook_path)[1] == ("E-1", None)
