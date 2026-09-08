import uuid
from pathlib import Path

from openpyxl import load_workbook
from rdflib import Graph

from infobim.context.ontology import (
    EntityContextContract,
    EntityFacadeContract,
    FacadeFieldContract,
    InfoBIMOntologyAdapter,
)
from infobim.context.workbook import EntityContextWorkbookAdapter


NS_TTL = """
@prefix ibim: <https://infobim.org/ontology/ns#> .
@prefix ifcowl: <https://standards.buildingsmart.org/IFC/DEV/IFC4/FINAL/OWL#> .

ibim:IfcWorkSchedule
    ibim:assignsRelatedClass ifcowl:IfcTask , ifcowl:IfcTaskTime .
"""

VIEW_TTL = """
@prefix ibim: <https://infobim.org/ontology/ns#> .
@prefix ifcowl: <https://standards.buildingsmart.org/IFC/DEV/IFC4/FINAL/OWL#> .
@prefix schema: <https://schema.org/> .
@prefix xsd: <http://www.w3.org/2001/XMLSchema#> .

ibim:WorkScheduleFacade
    ibim:targetsClass ibim:IfcWorkSchedule ;
    ibim:hasFacadeField ibim:ScheduleGlobalId , ibim:ScheduleName , ibim:ScheduleType .
ibim:ScheduleGlobalId schema:identifier "GlobalId" ; ibim:fieldOrder 10 ; ibim:fieldDatatype xsd:string ; ibim:isIdentifierField true .
ibim:ScheduleName schema:identifier "Name" ; ibim:fieldOrder 20 ; ibim:fieldDatatype xsd:string .
ibim:ScheduleType schema:identifier "PredefinedType" ; ibim:fieldOrder 30 ; ibim:fieldDatatype xsd:string .

ibim:TaskFacade
    ibim:targetsClass ifcowl:IfcTask ;
    ibim:hasFacadeField ibim:TaskGlobalId , ibim:TaskName , ibim:TaskTimeRef .
ibim:TaskGlobalId schema:identifier "GlobalId" ; ibim:fieldOrder 10 ; ibim:fieldDatatype xsd:string .
ibim:TaskName schema:identifier "Name" ; ibim:fieldOrder 20 ; ibim:fieldDatatype xsd:string .
ibim:TaskTimeRef schema:identifier "TaskTime" ; ibim:fieldOrder 30 ; ibim:fieldDatatype xsd:string .

ibim:TaskTimeFacade
    ibim:targetsClass ifcowl:IfcTaskTime ;
    ibim:hasFacadeField ibim:ScheduleStart , ibim:ScheduleFinish , ibim:ActualStart , ibim:ActualFinish .
ibim:ScheduleStart schema:identifier "ScheduleStart" ; ibim:fieldOrder 10 ; ibim:fieldDatatype xsd:dateTime .
ibim:ScheduleFinish schema:identifier "ScheduleFinish" ; ibim:fieldOrder 20 ; ibim:fieldDatatype xsd:dateTime .
ibim:ActualStart schema:identifier "ActualStart" ; ibim:fieldOrder 30 ; ibim:fieldDatatype xsd:dateTime .
ibim:ActualFinish schema:identifier "ActualFinish" ; ibim:fieldOrder 40 ; ibim:fieldDatatype xsd:dateTime .
"""


def test_ifc_work_schedule_discovers_related_classes(monkeypatch):
    ns_graph = Graph().parse(data=NS_TTL, format="turtle")
    view_graph = Graph().parse(data=VIEW_TTL, format="turtle")
    adapter = InfoBIMOntologyAdapter()
    monkeypatch.setattr(adapter, "_load_graphs", lambda: (ns_graph, view_graph))

    contract = adapter.resolve_entity_context("IfcWorkSchedule")

    assert contract.root.entity_name == "IfcWorkSchedule"
    assert [entity.entity_name for entity in contract.related] == [
        "IfcTask",
        "IfcTaskTime",
    ]


def test_ifc_guid_compression_matches_buildingsmart_example():
    source = uuid.UUID("f70dd363-bfe3-495d-84a0-2c02dcb7d4d2")
    assert (
        EntityContextWorkbookAdapter._compress_ifc_guid(source)
        == "3t3TDZl_D9NOIWB0BSjzJI"
    )


def test_schedule_context_materializes_one_workbook_with_three_sheets(
    tmp_path: Path,
):
    contract = EntityContextContract(
        root=EntityFacadeContract(
            entity_uri="https://infobim.org/ontology/ns#IfcWorkSchedule",
            entity_name="IfcWorkSchedule",
            facade_uri="https://infobim.org/ontology/ns#WorkScheduleFacade",
            fields=(
                FacadeFieldContract("GlobalId", order=10, identifier=True),
                FacadeFieldContract("Name", order=20),
                FacadeFieldContract("PredefinedType", order=30),
            ),
        ),
        related=(
            EntityFacadeContract(
                entity_uri="https://example.org/ifc#IfcTask",
                entity_name="IfcTask",
                facade_uri="https://infobim.org/ontology/ns#TaskFacade",
                fields=(
                    FacadeFieldContract("GlobalId", order=10),
                    FacadeFieldContract("Name", order=20),
                    FacadeFieldContract("TaskTime", order=30),
                ),
            ),
            EntityFacadeContract(
                entity_uri="https://example.org/ifc#IfcTaskTime",
                entity_name="IfcTaskTime",
                facade_uri="https://infobim.org/ontology/ns#TaskTimeFacade",
                fields=(
                    FacadeFieldContract("ScheduleStart", "dateTime", 10),
                    FacadeFieldContract("ScheduleFinish", "dateTime", 20),
                    FacadeFieldContract("ActualStart", "dateTime", 30),
                    FacadeFieldContract("ActualFinish", "dateTime", 40),
                ),
            ),
        ),
    )

    dataset_path = tmp_path / "dataset"
    dataset_path.mkdir()
    artifact = EntityContextWorkbookAdapter().generate(
        dataset_path=dataset_path,
        contract=contract,
        instance_name="Cronograma Executivo",
    )

    workbook = load_workbook(artifact.workbook_path)
    assert workbook.sheetnames == [
        "IfcWorkSchedule",
        "IfcTask",
        "IfcTaskTime",
    ]

    schedule = workbook["IfcWorkSchedule"]
    headers = [cell.value for cell in schedule[1]]
    row = dict(zip(headers, [cell.value for cell in schedule[2]]))
    assert row["Name"] == "Cronograma Executivo"
    assert row["PredefinedType"] == "PLANNED"
    assert isinstance(row["GlobalId"], str)
    assert len(row["GlobalId"]) == 22
    assert row["GlobalId"][0] in "0123"

    assert [cell.value for cell in workbook["IfcTask"][1]] == [
        "GlobalId",
        "Name",
        "TaskTime",
    ]
    assert [cell.value for cell in workbook["IfcTaskTime"][1]] == [
        "ScheduleStart",
        "ScheduleFinish",
        "ActualStart",
        "ActualFinish",
    ]

    validations = list(schedule.data_validations.dataValidation)
    assert len(validations) == 1
    assert (
        "ACTUAL,BASELINE,PLANNED,USERDEFINED,NOTDEFINED"
        in validations[0].formula1
    )
    workbook.close()
