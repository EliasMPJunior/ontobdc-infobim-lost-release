"""`infobim 4d --task` writes where the container says, and only there.

One task is three rows across three sheets, tied by GlobalIds. Typing those
by hand is where a schedule loses data silently: a wrong id raises nothing,
the task just stops appearing in the Gantt. These tests pin that the command
generates the ids, appends without disturbing what is already in the
workbook, and takes the file and sheet names from the container's own
datapackage rather than going looking for a spreadsheet.
"""
from __future__ import annotations

import datetime as dt
import importlib
import json
from pathlib import Path
from typing import Any, Dict, List

import pytest
from openpyxl import Workbook, load_workbook

from infobim.cli.adapter.help import (
    build_command_table,
    build_domain_help_content,
    discover_logical_components,
)
screen_module = importlib.import_module("infobim.4d.adapter.screen")
workbook_module = importlib.import_module("infobim.4d.adapter.workbook")
task_command_module = importlib.import_module("infobim.4d.plugin.command.task")

record_progress = screen_module.record_progress
record_task = screen_module.record_task
render = screen_module.render
ScheduleWorkbookAdapter = workbook_module.ScheduleWorkbookAdapter
ScheduleWorkbookError = workbook_module.ScheduleWorkbookError
parse_date = workbook_module.parse_date
parse_percent = workbook_module.parse_percent
wbs_sort_key = workbook_module.wbs_sort_key
FourDTaskCommand = task_command_module.FourDTaskCommand

WORKBOOK_RELATIVE_PATH: str = "../payload/document/cronograma.xlsx"


def build_container(root: Path, *, task_time_columns: List[str] = None) -> Path:
    """A container shaped the way `entity_workbook` writes one."""
    marker: Path = root / ".__ontobdc__"
    marker.mkdir(parents=True)
    document: Path = root / "payload" / "document"
    document.mkdir(parents=True)

    sheets: Dict[str, Any] = {
        "IfcWorkSchedule": (
            ["GlobalId", "Name", "PredefinedType"],
            [["WS-1", "Fundacao", "PLANNED"]],
        ),
        "IfcTask": (
            ["GlobalId", "Identification", "Name", "TaskTime"],
            [["T-1", "1.1", "Escavacao", "TT-1"]],
        ),
        "IfcTaskTime": (
            task_time_columns or ["GlobalId", "ScheduleStart", "ScheduleFinish"],
            [["TT-1", dt.datetime(2026, 8, 18), dt.datetime(2026, 8, 20)]],
        ),
        "IfcRelSequence": (
            ["GlobalId", "RelatingProcess", "RelatedProcess", "SequenceType"],
            [],
        ),
    }

    workbook: Workbook = Workbook()
    workbook.remove(workbook.active)
    resources: List[Dict[str, Any]] = []
    name: str
    for name, (header, rows) in sheets.items():
        worksheet = workbook.create_sheet(name)
        worksheet.append(header)
        for row in rows:
            worksheet.append(row)
        resources.append({
            "name": name.lower(),
            "path": WORKBOOK_RELATIVE_PATH,
            "format": "xlsx",
            "dialect": {"excel": {"sheet": name}},
            "entityUri": f"https://infobim.org/ontology/ns#{name}",
            "entityIdentifier": name,
        })
    workbook.save(document / "cronograma.xlsx")
    workbook.close()
    (marker / "datapackage.json").write_text(
        json.dumps({"resources": resources}), encoding="utf-8"
    )
    return root


@pytest.fixture
def container(tmp_path: Path) -> Path:
    return build_container(tmp_path / "container")


@pytest.fixture
def adapter(container: Path) -> ScheduleWorkbookAdapter:
    return ScheduleWorkbookAdapter(container)


# ------------------------------------------------------------- the mapping

def test_the_workbook_comes_from_the_datapackage(
    adapter: ScheduleWorkbookAdapter,
    container: Path,
) -> None:
    """Not from a search: the Surface reads the mapped file, so must this."""
    assert adapter.workbook_path == (
        container / "payload" / "document" / "cronograma.xlsx"
    )
    assert adapter.sheet_names["IfcTask"] == "IfcTask"
    assert adapter.sheet_names["IfcTaskTime"] == "IfcTaskTime"


def test_a_decoy_workbook_elsewhere_is_ignored(container: Path) -> None:
    decoy: Path = container / "outra.xlsx"
    workbook: Workbook = Workbook()
    for name in ("IfcWorkSchedule", "IfcTask", "IfcTaskTime", "IfcRelSequence"):
        workbook.create_sheet(name).append(["GlobalId"])
    workbook.remove(workbook["Sheet"])
    workbook.save(decoy)
    workbook.close()

    assert ScheduleWorkbookAdapter(container).workbook_path != decoy


def test_a_container_without_a_schedule_says_so(tmp_path: Path) -> None:
    bare: Path = tmp_path / "bare"
    (bare / ".__ontobdc__").mkdir(parents=True)
    (bare / ".__ontobdc__" / "datapackage.json").write_text(
        json.dumps({"resources": []}), encoding="utf-8"
    )

    with pytest.raises(ScheduleWorkbookError, match="no schedule"):
        ScheduleWorkbookAdapter(bare)


def test_no_datapackage_at_all_says_so(tmp_path: Path) -> None:
    with pytest.raises(ScheduleWorkbookError, match="datapackage.json"):
        ScheduleWorkbookAdapter(tmp_path / "nowhere-near-a-container")


# ------------------------------------------------------------- the writing

def test_recording_a_task_wires_its_three_rows(
    adapter: ScheduleWorkbookAdapter,
) -> None:
    record_task(adapter, {
        "wbs": "1.2",
        "name": "Armacao",
        "start": "2026-08-21",
        "finish": "2026-08-22",
        "after": "T-1",
    })

    workbook: Workbook = load_workbook(str(adapter.workbook_path))
    try:
        tasks = adapter.records(workbook, "IfcTask")
        times = adapter.records(workbook, "IfcTaskTime")
        sequences = adapter.records(workbook, "IfcRelSequence")
    finally:
        workbook.close()

    added = next(row for row in tasks if row["Identification"] == "1.2")
    time_row = next(
        row for row in times if row["GlobalId"] == added["TaskTime"]
    )

    # The join is the whole point: an id that does not resolve makes the task
    # disappear from the Gantt without any error anywhere.
    assert time_row["ScheduleStart"] == dt.datetime(2026, 8, 21)
    assert added["GlobalId"] not in {"T-1"}
    assert sequences[0]["RelatingProcess"] == "T-1"
    assert sequences[0]["RelatedProcess"] == added["GlobalId"]
    assert sequences[0]["SequenceType"] == "FINISH_START"


def test_existing_rows_are_kept(adapter: ScheduleWorkbookAdapter) -> None:
    record_task(adapter, {
        "wbs": "1.2", "name": "Armacao",
        "start": "2026-08-21", "finish": "2026-08-22",
    })

    assert [task["wbs"] for task in adapter.tasks()] == ["1.1", "1.2"]


def test_a_missing_column_is_created_rather_than_demanded(
    tmp_path: Path,
) -> None:
    """A schedule that only ever recorded the plan can start recording the
    actuals without anyone restructuring the workbook first."""
    container: Path = build_container(
        tmp_path / "planned-only",
        task_time_columns=["GlobalId", "ScheduleStart", "ScheduleFinish"],
    )
    adapter: ScheduleWorkbookAdapter = ScheduleWorkbookAdapter(container)

    record_task(adapter, {
        "wbs": "1.2", "name": "Armacao",
        "start": "2026-08-21", "finish": "2026-08-22",
        "actual_start": "2026-08-21", "percent": "40",
    })

    workbook: Workbook = load_workbook(str(adapter.workbook_path))
    try:
        header = [cell.value for cell in workbook["IfcTaskTime"][1]]
    finally:
        workbook.close()

    assert "ActualStart" in header
    assert "PercentComplete" in header


def test_columns_are_matched_by_header_not_by_position(
    tmp_path: Path,
) -> None:
    container: Path = build_container(
        tmp_path / "shuffled",
        task_time_columns=["ScheduleFinish", "GlobalId", "ScheduleStart"],
    )
    adapter: ScheduleWorkbookAdapter = ScheduleWorkbookAdapter(container)

    record_task(adapter, {
        "wbs": "1.2", "name": "Armacao",
        "start": "2026-08-21", "finish": "2026-08-22",
    })

    added = next(task for task in adapter.tasks() if task["wbs"] == "1.2")
    assert added["start"] == "2026-08-21"
    assert added["finish"] == "2026-08-22"


def test_a_duplicate_wbs_is_refused(adapter: ScheduleWorkbookAdapter) -> None:
    with pytest.raises(ValueError, match="1.1"):
        record_task(adapter, {
            "wbs": "1.1", "name": "Outra",
            "start": "2026-08-21", "finish": "2026-08-22",
        })


def test_a_finish_before_its_start_is_refused(
    adapter: ScheduleWorkbookAdapter,
) -> None:
    with pytest.raises(ValueError, match="anterior"):
        record_task(adapter, {
            "wbs": "1.9", "name": "Invertida",
            "start": "2026-08-29", "finish": "2026-08-20",
        })


def test_completion_records_a_finish_date(
    adapter: ScheduleWorkbookAdapter,
) -> None:
    """100% with no date cannot be compared against the plan, which is the
    only reason the actual columns exist."""
    record_progress(adapter, {"id": "T-1", "percent": "100"})

    done = next(task for task in adapter.tasks() if task["wbs"] == "1.1")
    assert done["percent"] == "100"
    assert done["actual_finish"] == dt.date.today().strftime("%Y-%m-%d")


def test_progress_starts_the_clock_once(
    adapter: ScheduleWorkbookAdapter,
) -> None:
    record_progress(adapter, {"id": "T-1", "percent": "30"})
    first = next(task for task in adapter.tasks() if task["wbs"] == "1.1")

    record_progress(adapter, {"id": "T-1", "percent": "60"})
    second = next(task for task in adapter.tasks() if task["wbs"] == "1.1")

    assert first["actual_start"] == second["actual_start"] == "2026-08-18"


# ----------------------------------------------------------- small parsers

@pytest.mark.parametrize(
    "written,expected",
    [
        ("2026-08-25", dt.datetime(2026, 8, 25)),
        ("25/08/2026", dt.datetime(2026, 8, 25)),
        ("25-08-2026", dt.datetime(2026, 8, 25)),
    ],
)
def test_dates_are_read_the_way_people_write_them(
    written: str,
    expected: dt.datetime,
) -> None:
    assert parse_date(written) == expected


def test_an_unreadable_date_is_reported() -> None:
    with pytest.raises(ValueError, match="não reconhecida"):
        parse_date("ontem")


@pytest.mark.parametrize(
    "written,expected", [("40", 40), ("40%", 40), ("40,5", 40), ("", None)]
)
def test_percentages_are_read_leniently(written: str, expected: Any) -> None:
    assert parse_percent(written) == expected


def test_percentages_are_clamped() -> None:
    assert parse_percent("150") == 100
    assert parse_percent("-10") == 0


def test_wbs_sorts_the_way_a_schedule_reads() -> None:
    ordered = sorted(["1.10", "1.9", "2.1", "1.2"], key=wbs_sort_key)
    assert ordered == ["1.2", "1.9", "1.10", "2.1"]


# --------------------------------------------------------------- the screen

def test_the_screen_lists_what_is_in_the_workbook(
    adapter: ScheduleWorkbookAdapter,
) -> None:
    page: str = render(adapter)

    assert "Escavacao" in page
    assert str(adapter.workbook_path) in page
    assert 'action="/task"' in page
    assert 'action="/progress"' in page


def test_a_late_finish_is_marked(adapter: ScheduleWorkbookAdapter) -> None:
    record_task(adapter, {
        "wbs": "1.2", "name": "Atrasada",
        "start": "2026-08-21", "finish": "2026-08-22",
        "actual_finish": "2026-08-27",
    })

    assert "⚠" in render(adapter)


def test_the_screen_escapes_what_it_shows(
    adapter: ScheduleWorkbookAdapter,
) -> None:
    record_task(adapter, {
        "wbs": "1.2", "name": '<script>alert("x")</script>',
        "start": "2026-08-21", "finish": "2026-08-22",
    })

    page: str = render(adapter)
    assert "<script>alert" not in page
    assert "&lt;script&gt;" in page


# -------------------------------------------------------------- the command

@pytest.mark.parametrize(
    "arguments",
    [
        ["4d"],
        ["4d", "--task"],
        ["4d", "--task", "--no-browser"],
        ["4d", "--container", "/tmp"],
        ["4d", "--task", "--container", "/tmp", "--port", "8800"],
    ],
)
def test_valid_invocations_are_accepted(arguments: List[str]) -> None:
    assert FourDTaskCommand.accepts(arguments)


@pytest.mark.parametrize(
    "arguments",
    [
        ["view"],
        ["4d", "--help"],
        ["4d", "--container"],
        ["4d", "--container", "--port"],
        ["4d", "--port", "oitocentos"],
        ["4d", "--task", "--task"],
        ["4d", "--unknown"],
    ],
)
def test_invalid_invocations_fall_through(arguments: List[str]) -> None:
    assert not FourDTaskCommand.accepts(arguments)


def test_the_domain_is_discovered() -> None:
    assert "4d" in discover_logical_components()


def test_the_command_appears_in_the_root_table() -> None:
    """Discovery is by convention here, so this is the only thing that says
    the command is reachable at all."""
    table = build_command_table(["4d"])
    assert "4d-task" in table
    assert table["4d-task"]["id"] == "4d_task"


def test_the_domain_help_names_the_action() -> None:
    """`build_domain_help_content` renders `arguments[0]` only, so the first
    entry has to be the one that says what the command does."""
    content = build_domain_help_content("4d")
    assert content["Usage"]["4d_task"] == "infobim 4d --task"
    assert "--task" in content["Options"]
