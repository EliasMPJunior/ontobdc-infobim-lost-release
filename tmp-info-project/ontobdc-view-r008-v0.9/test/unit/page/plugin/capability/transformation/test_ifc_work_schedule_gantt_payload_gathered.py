from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook

from ontobdc_view.page.adapter.context import PageDataContextAdapter
from ontobdc_view.page.plugin.capability.transformation.ifc_work_schedule_gantt_payload_gathered import (
    IfcWorkScheduleGanttPayloadGatheredCapability,
)
from ontobdc_view.surface.plugin.capability.transformation.data_gathered import (
    DataGatheredCapability,
)


IBIM = "https://infobim.org/ontology/ns#"


class FakeCliContext:
    raw_args: List[str] = []
    unprocessed_args: List[str] = []
    is_capability_targeted = False
    target_capability_id = None
    root_path = ""
    language = "en"

    def __init__(self, values: Dict[str, Any]) -> None:
        self.values = dict(values)

    def has_parameter(self, key: str) -> bool:
        return key in self.values

    def get_parameter_value(self, key: str) -> Any:
        return self.values.get(key)

    def set_parameter_value(self, key: str, value: Any) -> None:
        self.values[key] = value

    def delete_parameter(self, key: str) -> None:
        self.values.pop(key, None)

    def clear_parameters(self, keys: List[str]) -> None:
        for key in keys:
            self.delete_parameter(key)

    def reload(self) -> None:
        pass


def _schedule(element_uri: str) -> Dict[str, Any]:
    return {
        "@id": element_uri,
        "@type": [IBIM + "IfcWorkSchedule"],
        "http://purl.org/dc/terms/identifier": [{"@value": "schedule-1"}],
    }


def _write_data_gathered(parent: FakeCliContext, nodes: List[Dict[str, Any]]) -> None:
    state_path = DataGatheredCapability.state_path(parent)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(
        json.dumps(nodes, ensure_ascii=False),
        encoding="utf-8",
    )


def _page_context(
    parent: FakeCliContext,
    *,
    element_uri: str,
    schedule: Dict[str, Any],
    payload: Dict[str, Any],
) -> PageDataContextAdapter:
    return PageDataContextAdapter(
        parent,
        element_uri=element_uri,
        entity_uri=IBIM + "IfcWorkSchedule",
        source_node=schedule,
        payload=payload,
    )


def test_gantt_payload_preserves_detail_nodes_already_in_data_gathered(
    tmp_path: Path,
) -> None:
    element_uri = "https://example.test/schedule_dataset/schedule-1"
    schedule = _schedule(element_uri)
    task = {
        "@id": "https://example.test/schedule_dataset/task-1",
        "@type": [IBIM + "IfcTask"],
        IBIM + "GlobalId": [{"@value": "task-1"}],
        IBIM + "Name": [{"@value": "Mobilization"}],
    }
    parent = FakeCliContext({"container_path": str(tmp_path)})
    _write_data_gathered(parent, [task, schedule])

    payload: Dict[str, Any] = {}
    context = _page_context(
        parent,
        element_uri=element_uri,
        schedule=schedule,
        payload=payload,
    )

    result = IfcWorkScheduleGanttPayloadGatheredCapability().execute(context)

    assert result == {"resulting_state": "__gantt_payload_gathered__"}
    assert payload["@id"] == element_uri
    assert payload["@graph"][0] == schedule
    assert task in payload["@graph"]
    assert payload["gantt_payload"] == {
        "projectId": element_uri,
        "elementId": "schedule-1",
        "entity": "IfcWorkSchedule",
        "scheduleUri": element_uri,
        "resourceName": "ifc_work_schedule",
        "datasetFolder": "schedule_dataset",
        "datapackagePath": ".__ontobdc__/datapackage.json",
        "linksetPath": ".__ontobdc__/linkset/ns.ttl",
        "viewLinksetPath": ".__ontobdc__/linkset/view.ttl",
        "roCratePath": ".__ontobdc__/ro-crate-metadata.json",
    }
    assert payload["gantt_script_names"][-1] == "gantt_tab_mount"


def test_gantt_payload_materializes_task_time_and_sequence_sheets(
    tmp_path: Path,
) -> None:
    element_uri = "https://example.test/schedule_dataset/schedule-1"
    schedule = _schedule(element_uri)
    parent = FakeCliContext({"container_path": str(tmp_path)})
    _write_data_gathered(parent, [schedule])

    dataset_path = tmp_path / "schedule_dataset"
    dataset_path.mkdir(parents=True)
    workbook_path = dataset_path / "schedule.xlsx"

    workbook = Workbook()
    work_schedule = workbook.active
    work_schedule.title = "IfcWorkSchedule"
    work_schedule.append(["GlobalId", "Name"])
    work_schedule.append(["schedule-1", "Main Schedule"])

    task = workbook.create_sheet("IfcTask")
    task.append(["GlobalId", "Identification", "Name", "TaskTimeGlobalId"])
    task.append(["task-1", "1", "Mobilization", "time-1"])

    task_time = workbook.create_sheet("IfcTaskTime")
    task_time.append(
        ["GlobalId", "ScheduleStart", "ScheduleFinish", "Completion"]
    )
    task_time.append(["time-1", "2026-09-01", "2026-09-03", 50])

    sequence = workbook.create_sheet("IfcRelSequence")
    sequence.append(["GlobalId", "RelatingProcess", "RelatedProcess"])
    sequence.append(["sequence-1", "task-1", "task-2"])
    workbook.save(workbook_path)
    workbook.close()

    metadata_path = dataset_path / ".__ontobdc__"
    metadata_path.mkdir()
    (metadata_path / "datapackage.json").write_text(
        json.dumps(
            {
                "resources": [
                    {
                        "name": "ifc_work_schedule",
                        "entityIdentifier": "IfcWorkSchedule",
                        "entityUri": IBIM + "IfcWorkSchedule",
                        "path": "../schedule.xlsx",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    payload: Dict[str, Any] = {}
    context = _page_context(
        parent,
        element_uri=element_uri,
        schedule=schedule,
        payload=payload,
    )

    IfcWorkScheduleGanttPayloadGatheredCapability().execute(context)

    graph = payload["@graph"]
    assert graph[0] == schedule

    task_node = next(node for node in graph if node.get("@type") == [IBIM + "IfcTask"])
    assert task_node[IBIM + "GlobalId"] == [{"@value": "task-1"}]
    assert task_node[IBIM + "Name"] == [{"@value": "Mobilization"}]
    assert task_node[IBIM + "TaskTimeGlobalId"] == [{"@value": "time-1"}]

    time_node = next(
        node for node in graph if node.get("@type") == [IBIM + "IfcTaskTime"]
    )
    assert time_node[IBIM + "ScheduleStart"] == [{"@value": "2026-09-01"}]
    assert time_node[IBIM + "ScheduleFinish"] == [{"@value": "2026-09-03"}]
    assert time_node[IBIM + "Completion"] == [{"@value": "50"}]

    sequence_node = next(
        node for node in graph if node.get("@type") == [IBIM + "IfcRelSequence"]
    )
    assert sequence_node[IBIM + "RelatingProcess"] == [{"@value": "task-1"}]
    assert sequence_node[IBIM + "RelatedProcess"] == [{"@value": "task-2"}]
