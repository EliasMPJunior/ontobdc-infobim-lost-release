import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from ontobdc_view.page.adapter.container import (
    IFC_WORK_SCHEDULE_RUNTIME,
    WORK_STREAM_RUNTIME,
    chrome_controls_source,
)
from ontobdc_view.page.adapter.gantt_script import GanttScriptAdapter
from ontobdc_view.page.adapter.work_stream import WorkStreamScriptAdapter


def test_shared_inline_editor_swaps_pencil_for_save_and_cancel() -> None:
    for runtime in (WORK_STREAM_RUNTIME, IFC_WORK_SCHEDULE_RUNTIME):
        source = chrome_controls_source(runtime)
        assert "mountInlineEditor" in source
        assert 'iconButton("edit"' in source
        assert 'iconButton("save"' in source
        assert 'iconButton("cancel"' in source


def test_workstream_header_and_all_5w2h_dimensions_are_editable() -> None:
    adapter = WorkStreamScriptAdapter()
    graph = adapter.script_source("graph_reader")
    dimensions = adapter.script_source("dimension_card")

    assert 'saveWorkStreamField("Name"' in graph
    assert 'saveWorkStreamField("Description"' in graph
    for column in ("What", "Why", "Who", "Where", "When", "How", "HowMuch"):
        assert f'column: "{column}"' in dimensions
    assert "runtime.saveWorkStreamField(dimension.column" in dimensions
    assert "if (!value) continue" not in dimensions


def test_workstream_edits_are_persisted_to_the_connected_workbook() -> None:
    source = WorkStreamScriptAdapter().script_source("pyodide_runtime")

    assert "WORKBOOK_WRITE_SCRIPT" in source
    assert 'headers.get("GlobalId")' in source
    assert "workbook.save(workbook_path)" in source
    assert "syncMountedFilesystem" in source
    assert "liveWorkbookPath" in source
    assert "liveWorksheetName" in source


def test_workstream_write_script_updates_the_matching_xlsx_row(tmp_path: Path) -> None:
    source = WorkStreamScriptAdapter().script_source("pyodide_runtime")
    match = re.search(r"const WORKBOOK_WRITE_SCRIPT = `\n(.*?)\n`;", source, re.DOTALL)
    assert match is not None

    workbook_path = tmp_path / "workstream.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "WorkStream"
    sheet.append(["GlobalId", "Name", "What"])
    sheet.append(["ws-1", "Original", "Original scope"])
    workbook.save(workbook_path)
    workbook.close()

    namespace = {
        "work_stream_column_candidates_json": '["What"]',
        "workbook_path": str(workbook_path),
        "worksheet_name": "WorkStream",
        "work_stream_global_id": "ws-1",
        "work_stream_value": "Updated scope",
    }
    exec(match.group(1), namespace)

    persisted = load_workbook(workbook_path, read_only=True)
    try:
        assert persisted["WorkStream"]["C2"].value == "Updated scope"
    finally:
        persisted.close()


def test_schedule_title_and_description_are_editable_and_persisted() -> None:
    adapter = GanttScriptAdapter()
    persistence = adapter.script_source("pyodide_runtime")
    rendering = adapter.script_source("task_table_timeline")

    assert "writeScheduleMetadataToWorkbook" in persistence
    assert 'sheetName = "IfcWorkSchedule"' in persistence
    assert 'runtime.saveScheduleField = saveScheduleField' in persistence
    assert "renderScheduleHeader(schedule)" in rendering
    assert 'saveScheduleField("Name"' in rendering
    assert 'saveScheduleField("Description"' in rendering
